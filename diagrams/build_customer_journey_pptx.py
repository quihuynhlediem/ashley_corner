from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.oxml.ns import qn
from lxml import etree

PX=9525
def E(v): return Emu(int(v*PX))
NAVY=RGBColor(0x0A,0x23,0x78); INK=RGBColor(0x0A,0x1A,0x5C); MUTED=RGBColor(0x4A,0x5A,0x85)
BLUE=RGBColor(0x1F,0x5F,0xD6); SKY=RGBColor(0x33,0xA9,0xEE); RED=RGBColor(0xE3,0x26,0x2B)
WHITE=RGBColor(0xFF,0xFF,0xFF); LBLUE=RGBColor(0xE2,0xF3,0xFD); BORDER=RGBColor(0x9D,0xBD,0xEB); LINE=RGBColor(0x7C,0x8D,0xB5)
FONT="Noto Sans"

prs=Presentation(); prs.slide_width=E(1700); prs.slide_height=E(2000)
s=prs.slides.add_slide(prs.slide_layouts[6])
bg=s.background.fill; bg.solid(); bg.fore_color.rgb=WHITE

def runs(tf,lines,align=PP_ALIGN.CENTER):
    tf.word_wrap=True
    for i,(t,size,color,bold) in enumerate(lines):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.alignment=align
        r=p.add_run(); r.text=t; f=r.font; f.name=FONT; f.size=Pt(size); f.color.rgb=color; f.bold=bold

def shape(kind,x,y,w,h,title,subs=(),fill=WHITE,line=BORDER,lw=1.5,tcol=INK,scol=MUTED,dash=False,ts=15,ss=10.5):
    sh=s.shapes.add_shape(kind,E(x),E(y),E(w),E(h))
    sh.shadow.inherit=False
    sh.fill.solid(); sh.fill.fore_color.rgb=fill
    sh.line.color.rgb=line; sh.line.width=Pt(lw)
    if dash: sh.line.dash_style=MSO_LINE_DASH_STYLE.DASH
    tf=sh.text_frame; tf.vertical_anchor=MSO_ANCHOR.MIDDLE
    for m in ("margin_left","margin_right","margin_top","margin_bottom"): setattr(tf,m,E(8))
    runs(tf,[(title,ts,tcol,True)]+[(t,ss,scol,False) for t in subs])
    return sh

def badge(label,x,y):
    w=max(40,24+11*len(label))
    b=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,E(x-w/2),E(y-16),E(w),E(32))
    b.adjustments[0]=0.5; b.shadow.inherit=False
    b.fill.solid(); b.fill.fore_color.rgb=WHITE; b.line.color.rgb=NAVY; b.line.width=Pt(1.25)
    tf=b.text_frame; tf.vertical_anchor=MSO_ANCHOR.MIDDLE
    for m in ("margin_left","margin_right","margin_top","margin_bottom"): setattr(tf,m,0)
    runs(tf,[(label,9,NAVY,True)])

def text(x,y,w,h,t,size=10,color=MUTED,bold=False,align=PP_ALIGN.CENTER,bgfill=True):
    tb=s.shapes.add_textbox(E(x),E(y),E(w),E(h))
    if bgfill: tb.fill.solid(); tb.fill.fore_color.rgb=WHITE
    tf=tb.text_frame; tf.vertical_anchor=MSO_ANCHOR.MIDDLE
    for m in ("margin_left","margin_right","margin_top","margin_bottom"): setattr(tf,m,E(2))
    lines=t.split("\n"); runs(tf,[(l,size,color,bold) for l in lines],align)

def path(pts,dashed=False,arrow=True):
    n=len(pts)-1
    for i in range(n):
        (x1,y1),(x2,y2)=pts[i],pts[i+1]
        c=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,E(x1),E(y1),E(x2),E(y2))
        c.line.color.rgb=LINE; c.line.width=Pt(1.5)
        if dashed: c.line.dash_style=MSO_LINE_DASH_STYLE.DASH
        if arrow and i==n-1:
            ln=c.line._get_or_add_ln()
            te=etree.SubElement(ln,qn("a:tailEnd")); te.set("type","triangle"); te.set("w","med"); te.set("len","med")

# Lines first (behind shapes)
path([(355,240),(420,240)])
path([(742,240),(810,240)])
path([(1162,240),(1162,100),(1460,100),(1460,168)])
path([(988,324),(988,437)])
path([(988,590),(988,674)])
path([(812,763),(687,763)])
path([(330,760),(295,760),(295,515),(808,515)])
path([(988,851),(988,940)])
path([(1162,1029),(1207,1029),(1207,218),(1281,218)])
path([(812,1029),(540,1029),(540,1171)])
path([(718,1250),(808,1250)])
path([(1165,1252),(1232,1252),(1232,268),(1281,268)])
path([(1460,312),(1460,438)])
path([(1638,515),(1672,515),(1672,1400),(228,1400),(228,1433)])
path([(423,1540),(506,1540)])
path([(858,1540),(943,1540)])
path([(684,1628),(684,1720)])
path([(507,1797),(228,1797),(228,1647)])
path([(1114,1611),(1114,1720)],dashed=True)

# Header
text(38,30,700,60,"Customer journey",28,NAVY,True,PP_ALIGN.LEFT,False)
r=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,E(38),E(94),E(185),E(5)); r.fill.solid(); r.fill.fore_color.rgb=RED; r.line.fill.background(); r.shadow.inherit=False

T=MSO_SHAPE.FLOWCHART_TERMINATOR; R=MSO_SHAPE.RECTANGLE; D=MSO_SHAPE.DIAMOND
shape(T,35,170,320,142,"Checkout Interface",["Any website or application","(Shopee, CellphoneS...)"],LBLUE,SKY,2,NAVY,BLUE); badge("1",318,168)
shape(R,422,170,320,142,'Taps "Pay later"',["Pre-screened by Hong Leong Bank"]); badge("1",706,168)
shape(D,812,157,350,167,"Limit covers item price?",[],NAVY,NAVY,1.5,WHITE,WHITE,ts=11.5); badge("4",1082,196)
shape(R,1283,170,355,142,"Picks plan",["Tenor, monthly pay,","total cost shown"]); badge("6",1600,168)
shape(R,1283,440,355,150,"Gets item",["Hong Leong Bank pays merchant,","loan reported to CIC"]); badge("1 & 5",1581,438)
shape(R,810,440,355,150,"Shares data, ID",["One-tap consent,","chip ID + selfie"]); badge("3 & 5",1108,438)
shape(R,330,685,355,150,'"Data-for-limit" prompt',["Unlocks higher limit with more data"])
shape(D,812,676,350,175,"Instant decision?",[],NAVY,NAVY,1.5,WHITE,WHITE,ts=11.5); badge("2",1070,720)
shape(D,812,942,350,175,"Accept offer?",[],NAVY,NAVY,1.5,WHITE,WHITE,ts=11.5); badge("2",1074,982)
shape(R,363,1173,355,152,"Request re-assessment",["Credit request will be reviewed by banker"])
shape(R,810,1173,355,160,"Banker reviews",["Takes more than 30 min","(manual review by banker)"]); badge("3",1128,1172)
shape(R,35,1435,388,210,"Pays back in instalments",["Personal users: fixed monthly","Gig workers: % of daily pay","Merchants: % of daily sales"]); badge("7 & 8",365,1434)
shape(D,508,1453,350,175,"Hardship or return?",[],NAVY,NAVY,1.5,WHITE,WHITE,ts=11.5); badge("9",786,1492)
shape(T,945,1468,338,143,"Limit grows",["After on-time pay"],LBLUE,SKY,2,NAVY,BLUE); badge("4",1246,1468)
shape(R,507,1722,355,152,"Gets help",["Pause, extension","or refund"]); badge("9",824,1721)
shape(R,945,1722,338,152,"Optional offer",["Hong Leong Bank card or loan,","customer's choice"],dash=True); badge("11",1243,1721)

# Labels
text(1298,100,50,24,"Yes")
text(968,368,40,24,"No")
text(720,750,62,24,"declined")
text(820,862,336,42,"Approved\n(full price or lower limit (for low-confident case))",9.5)
text(1186,956,42,24,"Yes")
text(762,1017,40,24,"No")
text(1172,1262,90,24,"if approved",bgfill=False)
text(882,1528,40,24,"No")
text(664,1660,40,24,"Yes")
text(1048,1652,132,24,"strong history only")
text(450,1930,800,36,"↻ Every repayment helps retrain the decision engine (10)",12,NAVY,True,bgfill=False)

import os
out=os.path.join(os.path.dirname(os.path.abspath(__file__)),"customer-journey-canva.pptx")
prs.save(out); print(out)
